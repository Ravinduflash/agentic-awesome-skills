"""
Personal Finance Report Generator
===================================
Generates a Word document (.docx) covering cryptocurrency exchange
tiered earning rate principles, with a focus on MEXC Earn subscriptions,
and alternative yield-maximization strategies.

Usage:
    python reports/generate_finance_report.py
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUTPUT_PATH = os.path.join(os.path.dirname(__file__), "personal_finance_report.docx")


def set_cell_bg(cell, hex_color: str):
    """Set background colour of a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    return p


def add_paragraph(doc: Document, text: str, bold: bool = False, italic: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def add_bullet(doc: Document, text: str, bold_prefix: str = ""):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run_prefix = p.add_run(bold_prefix)
        run_prefix.bold = True
    p.add_run(text)
    return p


def build_document() -> Document:
    doc = Document()

    # ── Title ──────────────────────────────────────────────────────────────
    title = doc.add_heading("Personal Finance Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        "Cryptocurrency Exchange Tiered Earning Rates & Yield-Maximisation Strategies"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True

    doc.add_paragraph()  # spacer

    # ── Section 1 – Universal Principle ────────────────────────────────────
    add_heading(doc, "1. Universal Exchange Principle: Tiered Rate Limits", 1)

    add_paragraph(
        doc,
        "Every major cryptocurrency exchange operates on the same fundamental principle "
        "when it comes to tiered earning products: you cannot bypass the tier limits for "
        "a specific cryptocurrency by opening multiple separate Earn subscriptions.",
    )
    add_paragraph(
        doc,
        "This applies uniformly across leading platforms including MEXC, OKX, Binance, "
        "Bybit, and others. The mechanism is account-wide, not per-subscription.",
    )

    # ── Section 2 – MEXC Earn ──────────────────────────────────────────────
    add_heading(doc, "2. MEXC Earn Subscription Tiers", 1)

    add_paragraph(
        doc,
        "Tiered APRs on MEXC are based on the aggregate amount of an asset held in that "
        "product across all subscriptions, not per individual order or Fixed Deposit (FD).",
    )

    # 2.1 Cumulative Calculation
    add_heading(doc, "2.1 Cumulative Calculation", 2)
    add_paragraph(
        doc,
        "The applicable APR tier is determined by the total principal in your account "
        "snapshots. Consider the following USDT example where Tier 1 is capped at "
        "300 USDT at 20% APR:",
    )

    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    headers = ["Scenario", "Subscriptions", "Effective APR"]
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_bg(hdr_cells[i], "4472C4")
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    row1 = table.rows[1].cells
    row1[0].text = "Single subscription"
    row1[1].text = "1 × 300 USDT"
    row1[2].text = "300 USDT at 20%"

    row2 = table.rows[2].cells
    row2[0].text = "Two subscriptions"
    row2[1].text = "2 × 300 USDT (600 total)"
    row2[2].text = "300 USDT at 20% + 300 USDT at 10%"

    doc.add_paragraph()  # spacer after table

    # 2.2 Tier 1 Limits
    add_heading(doc, "2.2 Tier 1 Limits", 2)
    add_paragraph(
        doc,
        "Higher rates are usually promotional or designated for new users. They are "
        "capped at small amounts to prevent large-scale exploitation of the best rates.",
    )

    # 2.3 Account-Wide Tracking
    add_heading(doc, "2.3 Account-Wide Tracking", 2)
    add_paragraph(
        doc,
        "The platform's system tracks the total principal in your account snapshots to "
        "determine the applicable tier for interest distribution. Opening additional "
        "subscriptions for the same asset does not reset or duplicate the Tier 1 "
        "allocation.",
    )

    # ── Section 3 – Alternative Strategies ─────────────────────────────────
    add_heading(doc, "3. Alternative Strategies to Maximise Yield", 1)

    add_paragraph(
        doc,
        "If you want to deploy more capital at Tier 1 rates, consider the following "
        "strategies:",
    )

    add_bullet(
        doc,
        " Instead of concentrating all funds into one coin (e.g., USDT), split "
        "your capital across different assets that each carry their own independent "
        "Tier 1 limits — for example USDC, ETH, or SOL.",
        "Diversify Assets:",
    )

    add_bullet(
        doc,
        " MEXC allows the creation of multiple sub-accounts for different trading "
        "strategies. Note, however, that Earn products may apply KYC-based limits "
        "that are tied to the primary account owner.",
        "Use Sub-Accounts:",
    )

    add_bullet(
        doc,
        ' MEXC frequently runs time-limited "Exclusive for New Users" or '
        '"Locked Savings" events that offer higher fixed rates for specific short terms '
        "(e.g., 2-day or 7-day terms) that do not follow the standard tiering system.",
        "Check Promotions:",
    )

    # ── Section 4 – Summary ─────────────────────────────────────────────────
    add_heading(doc, "4. Summary", 1)

    summary_rows = [
        ("Bypass tier limits via multiple subscriptions?", "No — tiers are cumulative account-wide"),
        ("Does this apply to all exchanges?", "Yes — universal principle across MEXC, OKX, Binance, etc."),
        ("Best way to access more Tier 1 capital?", "Diversify into other coins, use sub-accounts, watch promotions"),
    ]

    tbl = doc.add_table(rows=len(summary_rows) + 1, cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = tbl.rows[0].cells
    hdr[0].text = "Question"
    hdr[1].text = "Answer"
    for cell in hdr:
        cell.paragraphs[0].runs[0].bold = True
        set_cell_bg(cell, "4472C4")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, (q, a) in enumerate(summary_rows, start=1):
        tbl.rows[i].cells[0].text = q
        tbl.rows[i].cells[1].text = a
        if i % 2 == 0:
            set_cell_bg(tbl.rows[i].cells[0], "DCE6F1")
            set_cell_bg(tbl.rows[i].cells[1], "DCE6F1")

    doc.add_paragraph()
    add_paragraph(
        doc,
        "Note: Exchange product details change frequently. Always verify current rates, "
        "tier limits, and promotional terms directly on the exchange before making "
        "financial decisions.",
        italic=True,
    )

    return doc


def main():
    doc = build_document()
    doc.save(OUTPUT_PATH)
    print(f"Report saved to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
